import io
import os
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from api import index
from api.index import Acareacao, app, gerar_xlsx, usuario_atual

BASE = dict(
    data_reclamacao="2026-05-31",
    ticket=1000007,
    entrega=2000007,
    cliente="  FULANA   DE TAL ",
    status="Recebido",
    procedencia="Improcedente",
    motorista=" Motorista Teste ",
    embarcador="MAGAZINE LUIZA",
)


def linha(**kw):
    return {"id": 1, **Acareacao(**BASE).model_dump(), "criado_por": "teste", "criado_em": None, **kw}


def test_acareacao_limpa_espacos_e_aplica_padrao():
    a = Acareacao(**BASE)
    assert a.cliente == "FULANA DE TAL" and a.motorista == "Motorista Teste"
    assert a.acareacao == "Acareação Motorista" and a.detalhe == ""


@pytest.mark.parametrize(
    "campo,valor",
    [
        ("status", "Pendente"),
        ("procedencia", "procedente"),
        ("acareacao", "Email"),
        ("cliente", "   "),
        ("motorista", ""),
        ("ticket", 0),
        ("entrega", -5),
        ("ticket", 10**15),
        ("data_reclamacao", "0026-10-15"),
        ("data_reclamacao", "2025-12-31"),
        ("data_reclamacao", (date.today() + timedelta(days=2)).isoformat()),
    ],
)
def test_acareacao_rejeita_valor_invalido(campo, valor):
    with pytest.raises(ValidationError):
        Acareacao(**{**BASE, campo: valor})


def test_acareacao_remove_caracteres_de_controle_e_exporta():
    a = Acareacao(**{**BASE, "cliente": "BELTRANA\x00 TESTE", "detalhe": "linha\x0bquebra\nsegunda linha"})
    assert a.cliente == "BELTRANA TESTE" and a.detalhe == "linhaquebra\nsegunda linha"
    ws = load_workbook(io.BytesIO(gerar_xlsx([{"id": 1, **a.model_dump()}]))).active
    assert ws["J2"].value == "linhaquebra\nsegunda linha"


def test_gerar_xlsx_cabecalho_linhas_e_tipos():
    ws = load_workbook(io.BytesIO(gerar_xlsx([linha(), linha(id=2, cliente="CICLANA")])))["Acareações"]
    assert tuple(c.value for c in ws[1]) == index.CABECALHO
    assert ws["A2"].value.date() == date(2026, 5, 31) and ws["A2"].number_format == "DD/MM/YYYY"
    assert ws["B2"].value == 1000007 and ws["D2"].value == "FULANA DE TAL" and ws["D3"].value == "CICLANA"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:J3"


def test_gerar_xlsx_nao_vira_formula():
    ws = load_workbook(io.BytesIO(gerar_xlsx([linha(cliente='=HYPERLINK("http://x")')]))).active
    assert ws["D2"].data_type == "s" and ws["D2"].value == '=HYPERLINK("http://x")'


cliente = TestClient(app)
EU = "00000000-0000-0000-0000-000000000001"


class Resp:
    def __init__(self, status_code, corpo=None):
        self.status_code, self._corpo = status_code, corpo

    def json(self):
        return self._corpo


@pytest.fixture
def como():
    def _como(admin):
        app.dependency_overrides[usuario_atual] = lambda: {"id": EU, "usuario": "teste", "admin": admin, "criado_em": None}

    yield _como
    app.dependency_overrides.clear()


def test_sem_token_da_401():
    r = cliente.get("/api/acareacoes")
    assert r.status_code == 401 and r.json()["detail"] == "Faça login para continuar."


def test_token_invalido_da_401(monkeypatch):
    monkeypatch.setattr(index, "supabase", lambda *a, **k: Resp(401))
    assert cliente.get("/api/acareacoes", headers={"Authorization": "Bearer x"}).status_code == 401


def test_token_sem_perfil_da_403(monkeypatch):
    conta = {"id": EU, "email": "intruso@acareacoes.local", "app_metadata": {}}
    monkeypatch.setattr(index, "supabase", lambda *a, **k: Resp(200, conta))
    assert cliente.get("/api/acareacoes", headers={"Authorization": "Bearer x"}).status_code == 403


def test_usuario_comum_nao_gerencia_usuarios(como):
    como(admin=False)
    assert cliente.get("/api/usuarios").status_code == 403
    assert cliente.post("/api/usuarios", json={"usuario": "novo", "senha": "12345678"}).status_code == 403
    assert cliente.delete(f"/api/usuarios/{EU}").status_code == 403


def test_admin_nao_remove_a_si_mesmo(como):
    como(admin=True)
    r = cliente.delete(f"/api/usuarios/{EU}")
    assert r.status_code == 400 and "si mesmo" in r.json()["detail"]


def test_novo_usuario_valida_nome_e_senha(como):
    como(admin=True)
    assert cliente.post("/api/usuarios", json={"usuario": "A B", "senha": "12345678"}).status_code == 422
    assert cliente.post("/api/usuarios", json={"usuario": "novo", "senha": "123"}).status_code == 422


def test_registro_invalido_da_422_em_portugues(como):
    como(admin=False)
    r = cliente.post("/api/acareacoes", json={**BASE, "status": "X"})
    assert r.status_code == 422 and r.json()["detail"] == "Confira os campos: status"


def test_exportar_devolve_xlsx_dos_ids(como, monkeypatch):
    como(admin=False)

    class Conn:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def execute(self, sql, params):
            self.params = params
            return self

        def fetchall(self):
            return [linha()]

    conn = Conn()
    monkeypatch.setattr(index, "conectar", lambda: conn)
    r = cliente.post("/api/exportar", json={"ids": [1]})
    assert r.status_code == 200 and conn.params == ([1],)
    assert r.headers["content-disposition"].startswith('attachment; filename="acareacoes_')
    assert load_workbook(io.BytesIO(r.content)).active["D2"].value == "FULANA DE TAL"


@pytest.mark.skipif(not os.environ.get("DATABASE_URL_TESTE"), reason="sem banco de teste (DATABASE_URL_TESTE)")
def test_crud_no_banco(como, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", os.environ["DATABASE_URL_TESTE"])
    como(admin=False)
    criado = cliente.post("/api/acareacoes", json=BASE)
    assert criado.status_code == 201
    rid = criado.json()["id"]
    try:
        assert criado.json()["criado_por"] == "teste" and criado.json()["cliente"] == "FULANA DE TAL"
        editado = cliente.put(f"/api/acareacoes/{rid}", json={**BASE, "status": "Indenizado", "motorista": "MOTORISTA  TESTE"})
        assert editado.json()["status"] == "Indenizado" and editado.json()["motorista"] == "Motorista Teste"
        assert any(r["id"] == rid for r in cliente.get("/api/acareacoes").json())
    finally:
        assert cliente.delete(f"/api/acareacoes/{rid}").status_code == 204
        with index.conectar() as c:
            c.execute("delete from apelidos where campo = 'motorista' and chave = 'motoristateste'")
    assert cliente.put(f"/api/acareacoes/{rid}", json=BASE).status_code == 404
    assert cliente.delete(f"/api/acareacoes/{rid}").status_code == 404


class Apelidos:
    """Conexão falsa com a tabela apelidos: insere se a chave é nova, senão devolve o nome já gravado."""

    def __init__(self, tabela):
        self.tabela = tabela

    def execute(self, sql, params):
        campo, chave, nome = params
        self.linha = {"nome": self.tabela.setdefault((campo, chave), nome)}
        return self

    def fetchone(self):
        return self.linha


def test_chave_nome_ignora_acento_caixa_e_espacos():
    assert index.chave_nome(" Fulano  Graça ") == index.chave_nome("fulano graca") == "fulanograca"
    assert index.chave_nome("WEB CONTINENTAL") == index.chave_nome("webcontinental")


def test_padronizar_usa_de_para_e_registra_nome_novo():
    tabela = {("motorista", "apelidoantigo"): "Motorista Padrão"}
    c = Apelidos(tabela)
    dados = index.padronizar(c, {"motorista": "Apelido Antigo", "embarcador": "Loja Nova", "cliente": "X"})
    assert dados == {"motorista": "Motorista Padrão", "embarcador": "Loja Nova", "cliente": "X"}
    assert index.padronizar(c, {"motorista": "apelido antigo", "embarcador": "LOJA  NOVA"})["embarcador"] == "Loja Nova"
    assert tabela[("embarcador", "lojanova")] == "Loja Nova"
