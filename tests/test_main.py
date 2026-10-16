import csv
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

import main

CAB_JAN = ["TICKET", "ENTREGA", "CLIENTE", "ACAREAÇÃO MOTORISTA", "STATUS", "PROCEDÊNCIA", "DATA DA RECLAMAÇÃO",
           "MOTORISTA", "EMBARCADOR", "DATA DA RESOLUÇÃO", "VALOR", "DETALHE "]
CAB_NOVO = ["TICKET", "TICKET", "TICKET", "ACAREAÇÃO MOTORISTA", "STATUS", "PROCEDÊNCIA", "MOTORISTA", "EMBARCADOR",
            "DETALHE "]
REAL = Path(__file__).parent.parent / "ACAREAÇÕES HRAO.xlsx"


@pytest.fixture
def planilha(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    modelo = wb.create_sheet("MODELO")
    modelo.append(CAB_JAN)
    modelo.append([None, 1, "X", "Whats Motorista", "Aguardando", "Analizando"])
    jan = wb.create_sheet("JANEIRO")
    jan.append(CAB_JAN)
    jan.append([1000001.0, 2000001.0, "CLIENTE  UM ", "Acareação Motorista", "Recebido", "Improcedente",
                datetime(2026, 1, 2), "Motorista Alfa", "KABUM", datetime(2026, 3, 1), "R$ 100,00", None])
    jan.append([1000002.0, 2000002.0, "CLIENTE DOIS", "Acareação Motorista", "Aguardando", "Analizando", None,
                "Motorista Graca", "ELECTROLUX", None, None, "cliente sumiu"])
    mar = wb.create_sheet("MARÇO")
    mar.append(CAB_NOVO)
    mar.append([1000003.0, 2000003.0, "CLIENTE TRES", "Acareação Motorista", "Recebido", "procedente", "Motorista Graça",
                "MAGAZINE LUIZA", None])
    mar.append([1000004.0, 2000004.0, "CLIENTE QUATRO", None, None, None, "Motorista  Graça", "DAFITTI", None])
    ago = wb.create_sheet("AGOSTO")
    ago.append(CAB_NOVO)
    ago.append(["DIA 06"])
    ago.append([1000005.0, 2000005.0, "CLIENTE CINCO", "Acareação Motorista", "Recebido", "Improcedente", "Motorista Beta",
                "DAFITI", "falta interna"])
    ago.append(["Dia 07 "])
    ago.append([None, None, None, "Acareação Motorista"])
    ago.append([1000006.0, 2000006.0, "CLIENTE SEIS", "Acareação Motorista", "Devolvido", "Improcedente", "Motorista Beta",
                "DAFITI", None])
    res = wb.create_sheet("RESULTADO")
    res.append(["MOTORISTA", "IMPROCEDENTE"])
    res.append(["Alex", 5])
    caminho = tmp_path / "planilha.xlsx"
    wb.save(caminho)
    return caminho


class Conn:
    """Conexão falsa: guarda o que seria gravado."""

    def __init__(self, ja_tem=0):
        self.ja_tem, self.gravados, self.apelidos = ja_tem, [], []

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, sql, *a):
        return self

    def fetchone(self):
        return {"count": self.ja_tem}

    def cursor(self):
        return self

    def executemany(self, sql, linhas):
        (self.apelidos if "apelidos" in sql else self.gravados).extend(linhas)


def test_le_so_abas_de_mes_e_ignora_linhas_sem_ticket(planilha):
    assert [r["ticket"] for r in main.ler_planilha(planilha)] == [1000001, 1000002, 1000003, 1000004, 1000005, 1000006]


def test_layout_janeiro_preserva_valor_e_resolucao_no_detalhe(planilha):
    r = main.ler_planilha(planilha)[0]
    assert (r["data_reclamacao"], r["motorista"], r["embarcador"]) == (date(2026, 1, 2), "Motorista Alfa", "KABUM")
    assert r["cliente"] == "CLIENTE UM"
    assert r["detalhe"] == "Valor: R$ 100,00 · Resolução: 01/03/2026"


def test_datas_sem_dia_usam_dia_1_e_dia_nn_vale_para_as_linhas_abaixo(planilha):
    datas = [r["data_reclamacao"] for r in main.ler_planilha(planilha)]
    assert datas[1:] == [date(2026, 1, 1), date(2026, 3, 1), date(2026, 3, 1), date(2026, 8, 6), date(2026, 8, 7)]


def test_normaliza_listas_e_preenche_vazios(planilha):
    regs = main.ler_planilha(planilha)
    assert regs[1]["procedencia"] == "Analisando" and regs[1]["detalhe"] == "cliente sumiu"
    assert regs[2]["procedencia"] == "Procedente"
    assert (regs[3]["acareacao"], regs[3]["status"], regs[3]["procedencia"]) == (
        "Acareação Motorista", "Aguardando", "Aguardando")


def test_tudo_valido(planilha):
    validos, erros = main.validar(main.ler_planilha(planilha))
    assert erros == [] and len(validos) == 6 and "aba" not in validos[0]


def test_validar_aponta_aba_linha_e_campo():
    r = {"aba": "MAIO", "linha": 7, "data_reclamacao": date(2026, 5, 1), "ticket": 1, "entrega": 2, "cliente": "A",
         "acareacao": "Acareação Motorista", "status": "Sumido", "procedencia": "Aguardando", "motorista": "M",
         "embarcador": "E", "detalhe": ""}
    assert main.validar([r]) == ([], ["MAIO linha 7: status"])


def test_agrupa_grafias_parecidas_na_mais_frequente(planilha):
    g = {(x["campo"], x["original"]): x["sugerido"] for x in main.agrupar_nomes(main.ler_planilha(planilha))}
    assert g[("motorista", "Motorista Graca")] == g[("motorista", "Motorista Graça")] == "Motorista Graça"
    assert g[("embarcador", "DAFITTI")] == "DAFITI"
    assert g[("motorista", "Motorista Beta")] == "Motorista Beta"


def test_csv_ida_e_volta_ignora_sugestao_vazia(tmp_path):
    linhas = [
        {"campo": "motorista", "original": "Motorista Graca", "sugerido": "Motorista Graça", "ocorrencias": 1},
        {"campo": "motorista", "original": "Motorista Beta", "sugerido": "", "ocorrencias": 2},
    ]
    main.escrever_csv(linhas, tmp_path / "n.csv")
    assert main.ler_csv(tmp_path / "n.csv") == {("motorista", "Motorista Graca"): "Motorista Graça"}


def test_previa_nao_sobrescreve_nomes_csv_revisado(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "nomes.csv").write_text("revisado", encoding="utf-8")
    main.importar(planilha, aplicar=False)
    assert (tmp_path / "nomes.csv").read_text(encoding="utf-8") == "revisado"


def test_aplicar_sem_csv_para(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(SystemExit, match="Rode primeiro sem --aplicar"):
        main.importar(planilha, aplicar=True)


def test_aplicar_com_erro_nao_abre_o_banco(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    main.importar(planilha, aplicar=False)
    ruim = {**main.ler_planilha(planilha)[0], "status": "Sumido"}
    monkeypatch.setattr(main, "ler_planilha", lambda c: [ruim])
    monkeypatch.setattr(main, "conectar", lambda: pytest.fail("não deveria abrir o banco"))
    with pytest.raises(SystemExit, match="Nada foi gravado"):
        main.importar(planilha, aplicar=True)


def test_aplicar_recusa_tabela_com_registros(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    main.importar(planilha, aplicar=False)
    conn = Conn(ja_tem=3)
    monkeypatch.setattr(main, "conectar", lambda: conn)
    with pytest.raises(SystemExit, match="já tem registros"):
        main.importar(planilha, aplicar=True)
    assert conn.gravados == []


def test_aplicar_grava_tudo_com_nomes_unificados(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    main.importar(planilha, aplicar=False)
    assert (tmp_path / "nomes.csv").exists()
    conn = Conn()
    monkeypatch.setattr(main, "conectar", lambda: conn)
    main.importar(planilha, aplicar=True)
    assert len(conn.gravados) == 6
    assert {g["motorista"] for g in conn.gravados} == {"Motorista Alfa", "Motorista Graça", "Motorista Beta"}
    assert {g["embarcador"] for g in conn.gravados} >= {"DAFITI"} and "DAFITTI" not in {g["embarcador"] for g in conn.gravados}
    assert all(g["criado_por"] == "importacao" for g in conn.gravados)


def test_aplicar_grava_o_de_para_pela_chave_sem_acento_caixa_e_espacos(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    main.importar(planilha, aplicar=False)
    conn = Conn()
    monkeypatch.setattr(main, "conectar", lambda: conn)
    main.importar(planilha, aplicar=True)
    de_para = {(a["campo"], a["chave"]): a["nome"] for a in conn.apelidos}
    assert de_para[("motorista", "motoristagraca")] == "Motorista Graça"
    assert de_para[("embarcador", "dafitti")] == de_para[("embarcador", "dafiti")] == "DAFITI"
    assert de_para[("motorista", "motoristabeta")] == "Motorista Beta"


def test_aplicar_recusa_de_para_contraditorio(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    main.importar(planilha, aplicar=False)
    linhas = list(csv.DictReader(open("nomes.csv", encoding="utf-8-sig"), delimiter=";"))
    for l in linhas:  # duas grafias com a mesma chave apontando para nomes diferentes
        if l["original"] == "Motorista Graca":
            l["sugerido"] = "Outro Nome"
    main.escrever_csv(linhas)
    monkeypatch.setattr(main, "conectar", lambda: pytest.fail("não deveria abrir o banco"))
    with pytest.raises(SystemExit, match="motoristagraca"):
        main.importar(planilha, aplicar=True)


@pytest.mark.skipif(not REAL.exists(), reason="planilha real fora da pasta (dados pessoais, fora do git)")
def test_planilha_real_importa_os_956_registros():
    validos, erros = main.validar(main.ler_planilha(REAL))
    assert erros == [] and len(validos) == 956


def test_exigir_variaveis_explica_o_que_falta(monkeypatch):
    monkeypatch.setenv("SUPABASE_URL", "https://x.supabase.co")
    monkeypatch.setenv("SUPABASE_SERVICE_ROLE_KEY", "")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    with pytest.raises(SystemExit) as e:
        main.exigir("SUPABASE_URL", "SUPABASE_SERVICE_ROLE_KEY", "DATABASE_URL")
    assert "SUPABASE_SERVICE_ROLE_KEY, DATABASE_URL" in str(e.value) and "SUPABASE_URL," not in str(e.value)
    assert "invisíveis" in str(e.value)


def test_exigir_recusa_supabase_url_incompleta(monkeypatch):
    monkeypatch.setenv("SUPABASE_URL", "abcdefghij")
    with pytest.raises(SystemExit, match="https://abcdefghij.supabase.co"):
        main.exigir("SUPABASE_URL")
