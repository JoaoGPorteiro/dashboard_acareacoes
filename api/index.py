"""API do dashboard de acareações. Roda como função Python na Vercel (variável `app`)."""

import io
import os
import unicodedata
from datetime import date, timedelta
from typing import Literal
from uuid import UUID

import httpx
import psycopg
from fastapi import Depends, FastAPI, Header, HTTPException, Request, Response
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from psycopg.rows import dict_row
from pydantic import BaseModel, Field, field_validator

DOMINIO = "acareacoes.local"  # usuário "fulano" existe no Supabase Auth como fulano@acareacoes.local
PERFIS = {"admin", "usuario"}
INICIO = date(2026, 1, 1)  # primeiro mês do histórico; o "Tudo" do painel começa aqui

COLUNAS = (
    "data_reclamacao", "ticket", "entrega", "cliente", "acareacao",
    "status", "procedencia", "motorista", "embarcador", "detalhe",
)
CABECALHO = (
    "DATA DA RECLAMAÇÃO", "TICKET", "Nº DA ENTREGA", "CLIENTE", "ACAREAÇÃO MOTORISTA",
    "STATUS", "PROCEDÊNCIA", "MOTORISTA", "EMBARCADOR", "DETALHE",
)
LARGURAS = (14, 12, 14, 28, 22, 13, 14, 22, 22, 60)

# Nomes de coluna vêm só de COLUNAS (constante), nunca do usuário.
INSERIR = (
    f"insert into acareacoes ({', '.join(COLUNAS)}, criado_por) "
    f"values ({', '.join(f'%({c})s' for c in COLUNAS)}, %(criado_por)s) returning *"
)
ATUALIZAR = f"update acareacoes set {', '.join(f'{c} = %({c})s' for c in COLUNAS)} where id = %(id)s returning *"


class Acareacao(BaseModel):
    data_reclamacao: date
    ticket: int = Field(gt=0, lt=10**15)
    entrega: int = Field(gt=0, lt=10**15)
    cliente: str = Field(min_length=1, max_length=200)
    acareacao: Literal["Whatsapp", "Whats Motorista", "Acareação Motorista", "Recebido"] = "Acareação Motorista"
    status: Literal["Aguardando", "Recebido", "Indenizado", "Devolvido", "Parado"]
    procedencia: Literal["Procedente", "Improcedente", "Analisando", "Aguardando", "Sem retorno"]
    motorista: str = Field(min_length=1, max_length=200)
    embarcador: str = Field(min_length=1, max_length=200)
    detalhe: str = Field(default="", max_length=5000)

    @field_validator("cliente", "motorista", "embarcador", mode="before")
    @classmethod
    def _espacos(cls, v):
        return " ".join(ILLEGAL_CHARACTERS_RE.sub("", v).split()) if isinstance(v, str) else v

    @field_validator("detalhe", mode="before")
    @classmethod
    def _detalhe(cls, v):
        # caracteres de controle (ex.: colados do Word) quebram o .xlsx e o Postgres
        return "" if v is None else ILLEGAL_CHARACTERS_RE.sub("", v).strip() if isinstance(v, str) else v

    @field_validator("data_reclamacao")
    @classmethod
    def _data(cls, v: date) -> date:
        # +1 dia: o servidor roda em UTC, à frente de Brasília depois das 21h
        if not INICIO <= v <= date.today() + timedelta(days=1):
            raise ValueError("data fora do período válido")
        return v


def gerar_xlsx(linhas: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Acareações"
    ws.append(CABECALHO)
    for cel in ws[1]:
        cel.font = Font(bold=True)
    for linha in linhas:
        ws.append([linha[c] for c in COLUNAS])
        for cel in ws[ws.max_row]:
            if isinstance(cel.value, str) and cel.value.startswith("="):
                cel.data_type = "s"  # texto digitado não pode virar fórmula no Excel
    for cel in ws["A"][1:]:
        cel.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, largura in enumerate(LARGURAS):
        ws.column_dimensions[chr(65 + i)].width = largura
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class NovoUsuario(BaseModel):
    usuario: str = Field(pattern=r"^[a-z0-9._-]{3,40}$")
    senha: str = Field(min_length=8, max_length=72)
    admin: bool = False


class Exportacao(BaseModel):
    ids: list[int] = Field(max_length=50_000)


app = FastAPI(title="Acareações", docs_url=None, redoc_url=None, openapi_url=None)


@app.exception_handler(RequestValidationError)
async def _erro_validacao(_: Request, exc: RequestValidationError) -> JSONResponse:
    campos = sorted({str(e["loc"][-1]) for e in exc.errors()})
    return JSONResponse({"detail": "Confira os campos: " + ", ".join(campos)}, status_code=422)


def supabase(metodo: str, caminho: str, token: str | None = None, **kw) -> httpx.Response:
    chave = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    cabecalhos = {"apikey": chave, "Authorization": f"Bearer {token or chave}"}
    return httpx.request(metodo, os.environ["SUPABASE_URL"].rstrip("/") + caminho, headers=cabecalhos, timeout=10, **kw)


def _resumo(u: dict) -> dict:
    return {
        "id": u["id"],
        "usuario": u["email"].split("@")[0],
        "admin": u.get("app_metadata", {}).get("role") == "admin",
        "criado_em": u.get("created_at"),
    }


def usuario_atual(authorization: str = Header(default="")) -> dict:
    token = authorization.removeprefix("Bearer ").strip()
    if not token:
        raise HTTPException(401, "Faça login para continuar.")
    r = supabase("GET", "/auth/v1/user", token)
    if r.status_code != 200:
        raise HTTPException(401, "Sessão expirada. Entre novamente.")
    u = r.json()
    # Conta sem perfil (ex.: cadastro público ligado por engano no Supabase) não entra.
    if u.get("app_metadata", {}).get("role") not in PERFIS:
        raise HTTPException(403, "Usuário sem acesso ao dashboard.")
    return _resumo(u)


def so_admin(u: dict = Depends(usuario_atual)) -> dict:
    if not u["admin"]:
        raise HTTPException(403, "Apenas administradores podem fazer isso.")
    return u


def chave_nome(v: str) -> str:
    """Forma justa de comparar nomes: sem acento, sem caixa e sem espaços ("Fulano  Graça" == "fulano graca")."""
    return "".join(unicodedata.normalize("NFKD", v).encode("ascii", "ignore").decode().casefold().split())


def padronizar(c, dados: dict) -> dict:
    """Troca motorista e embarcador pelo nome padrão do de/para (tabela apelidos).

    Grafia nunca vista vira o padrão dela mesma, para as próximas variações caírem nela.
    """
    for campo in ("motorista", "embarcador"):
        dados[campo] = c.execute(
            "insert into apelidos (campo, chave, nome) values (%s, %s, %s) "
            "on conflict (campo, chave) do update set nome = apelidos.nome returning nome",
            (campo, chave_nome(dados[campo]), dados[campo]),
        ).fetchone()["nome"]
    return dados


def conectar() -> psycopg.Connection:
    # prepare_threshold=None: o pooler do Supabase (modo transação) não aceita prepared statements.
    return psycopg.connect(os.environ["DATABASE_URL"], prepare_threshold=None, row_factory=dict_row)


@app.get("/api/acareacoes")
def listar(_: dict = Depends(usuario_atual)) -> list[dict]:
    # shortcut: devolve tudo (~1.200/ano); paginar se passar de ~20 mil registros.
    with conectar() as c:
        return c.execute("select * from acareacoes order by data_reclamacao desc, id desc").fetchall()


@app.post("/api/acareacoes", status_code=201)
def criar(a: Acareacao, u: dict = Depends(usuario_atual)) -> dict:
    with conectar() as c:
        return c.execute(INSERIR, {**padronizar(c, a.model_dump()), "criado_por": u["usuario"]}).fetchone()


@app.put("/api/acareacoes/{rid}")
def editar(rid: int, a: Acareacao, _: dict = Depends(usuario_atual)) -> dict:
    with conectar() as c:
        linha = c.execute(ATUALIZAR, {**padronizar(c, a.model_dump()), "id": rid}).fetchone()
        if not linha:  # dentro da transação: desfaz o apelido gravado por padronizar
            raise HTTPException(404, "Registro não encontrado.")
    return linha


@app.delete("/api/acareacoes/{rid}", status_code=204)
def excluir(rid: int, _: dict = Depends(usuario_atual)) -> Response:
    with conectar() as c:
        if c.execute("delete from acareacoes where id = %s", (rid,)).rowcount == 0:
            raise HTTPException(404, "Registro não encontrado.")
    return Response(status_code=204)


@app.post("/api/exportar")
def exportar(e: Exportacao, _: dict = Depends(usuario_atual)) -> Response:
    with conectar() as c:
        linhas = c.execute(
            "select * from acareacoes where id = any(%s) order by data_reclamacao desc, id desc", (e.ids,)
        ).fetchall()
    return Response(
        gerar_xlsx(linhas),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="acareacoes_{date.today().isoformat()}.xlsx"'},
    )


def criar_usuario(novo: NovoUsuario) -> dict:
    r = supabase("POST", "/auth/v1/admin/users", json={
        "email": f"{novo.usuario}@{DOMINIO}",
        "password": novo.senha,
        "email_confirm": True,
        "app_metadata": {"role": "admin" if novo.admin else "usuario"},
    })
    if r.status_code >= 400:
        if "email_exists" in r.text or "already" in r.text:
            raise HTTPException(409, "Esse usuário já existe.")
        raise HTTPException(400, f"Não foi possível criar o usuário (Supabase {r.status_code}).")
    return _resumo(r.json())


@app.get("/api/usuarios")
def listar_usuarios(_: dict = Depends(so_admin)) -> list[dict]:
    r = supabase("GET", "/auth/v1/admin/users", params={"per_page": 1000})
    r.raise_for_status()
    return sorted((_resumo(u) for u in r.json()["users"]), key=lambda u: u["usuario"])


@app.post("/api/usuarios", status_code=201)
def novo_usuario(novo: NovoUsuario, _: dict = Depends(so_admin)) -> dict:
    return criar_usuario(novo)


@app.delete("/api/usuarios/{uid}", status_code=204)
def remover_usuario(uid: UUID, u: dict = Depends(so_admin)) -> Response:
    if str(uid) == u["id"]:
        raise HTTPException(400, "Você não pode remover a si mesmo.")
    r = supabase("DELETE", f"/auth/v1/admin/users/{uid}")
    if r.status_code == 404:
        raise HTTPException(404, "Usuário não encontrado.")
    r.raise_for_status()
    return Response(status_code=204)
